"""Read uploaded files (.csv / .xlsx / .xlsm / .docx) and write the colored Excel output."""
import csv
import io
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException

from app.core.config import PRINCIPAL_COLORS

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    Document = None

logger = logging.getLogger("file_io")

MU_ID_NAMES = ["MU_ID", "segment_id", "segment id", "id"]
SPEAKER_NAMES = ["Speaker", "speaker_name", "speaker name", "speaker_id", "speaker id"]
TEXT_NAMES = ["English Translation", "english_translation", "english", "text", "transcript"]
WORD_TEXT_NAMES = ["English Translation", "english_translation", "english translation", "english"]


# ---------------------------------------------------------------- helpers
def find_field(fieldnames: list[str], names: list[str]) -> str | None:
    normalized = {f.lower().strip(): f for f in fieldnames if f}
    for name in names:
        found = normalized.get(name.lower().strip())
        if found:
            return found
    return None


def _resolve_columns(fieldnames: list[str], source: str) -> tuple[str, str]:
    mu_id = find_field(fieldnames, MU_ID_NAMES)
    text = find_field(fieldnames, TEXT_NAMES)
    if not mu_id:
        raise HTTPException(400, f"{source} must contain a MU_ID or segment_id column.")
    if not text:
        raise HTTPException(400, f"{source} must contain an 'English Translation' column.")
    return mu_id, text


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _generated_id(position: int, taken: set[str]) -> str:
    """Plain running number (1, 2, 3 ...). Falls back to ROW_<n> only if the upload already uses that number."""
    candidate = str(position)
    return candidate if candidate not in taken else f"ROW_{position}"


def _fill_blank_ids(segments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Give every segment without an ID its row number, counting from 1. IDs from the file are kept."""
    taken = {s["segment_id"] for s in segments} - {""}
    for position, segment in enumerate(segments, start=1):
        if not segment["segment_id"]:
            segment["segment_id"] = _generated_id(position, taken)
    return segments


# ---------------------------------------------------------------- CSV
def read_csv_file(raw: bytes) -> list[dict[str, Any]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise HTTPException(400, "CSV is empty.")

    mu_id_field, text_field = _resolve_columns(reader.fieldnames or [], "CSV")

    return _fill_blank_ids(
        [
            {
                "segment_id": str(row.get(mu_id_field, "") or "").strip(),
                "text": str(row.get(text_field, "") or "").strip(),
                "original": dict(row),
            }
            for row in rows
        ]
    )


# ---------------------------------------------------------------- Excel
def read_excel_file(raw: bytes) -> list[dict[str, Any]]:
    if openpyxl is None:
        raise HTTPException(500, "openpyxl is not installed. Run: pip install openpyxl")

    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        rows_iter = workbook.active.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            raise HTTPException(400, "Excel file is empty.")

        fieldnames = [
            str(h).strip() if h is not None else f"column_{i + 1}"
            for i, h in enumerate(headers)
        ]
        mu_id_field, text_field = _resolve_columns(fieldnames, "Excel")
        mu_idx, text_idx = fieldnames.index(mu_id_field), fieldnames.index(text_field)

        segments = []
        for values in rows_iter:
            row = list(values) + [None] * (len(fieldnames) - len(values))
            segments.append(
                {
                    "segment_id": str(row[mu_idx] or "").strip(),
                    "text": str(row[text_idx] or "").strip(),
                    "original": {f: "" if row[i] is None else str(row[i]) for i, f in enumerate(fieldnames)},
                }
            )
        return _fill_blank_ids(segments)
    finally:
        workbook.close()


# ---------------------------------------------------------------- Word
def _word_segment(segment_id: str, text: str) -> dict[str, Any]:
    # Only MU_ID and English Translation are kept for Word files.
    return {
        "segment_id": segment_id,
        "text": text,
        "original": {"MU_ID": segment_id, "English Translation": text},
    }


def _number_word_segments(items: list[tuple[str, str]]) -> list[dict[str, Any]]:
    """items = [(id_from_file_or_blank, text)]. Blank IDs become 1, 2, 3 ... in reading order."""
    taken = {segment_id for segment_id, _ in items} - {""}
    return [
        _word_segment(segment_id or _generated_id(position, taken), text)
        for position, (segment_id, text) in enumerate(items, start=1)
    ]


def read_word_file(raw: bytes) -> list[dict[str, Any]]:
    """
    1. If any table has an 'English Translation' column -> use ONLY that column
       (+ MU_ID if present) from ALL such tables.
    2. Otherwise -> every non-empty paragraph is one segment.
    """
    if Document is None:
        raise HTTPException(500, "python-docx is not installed. Run: pip install python-docx")

    document = Document(io.BytesIO(raw))
    items: list[tuple[str, str]] = []  # (id from the file or "", text)
    found_table = False

    for table in document.tables:
        if not table.rows:
            continue
        headers = [_clean(c.text) for c in table.rows[0].cells]
        english_field = find_field(headers, WORD_TEXT_NAMES)
        if not english_field:
            continue

        found_table = True
        english_idx = headers.index(english_field)
        mu_field = find_field(headers, MU_ID_NAMES)
        mu_idx = headers.index(mu_field) if mu_field else None

        for row in table.rows[1:]:
            values = [_clean(c.text) for c in row.cells]
            if len(values) <= english_idx or not values[english_idx]:
                continue

            segment_id = values[mu_idx] if mu_idx is not None and mu_idx < len(values) else ""
            items.append((segment_id, values[english_idx]))

    if found_table:
        segments = _number_word_segments(items)
        if not segments:
            raise HTTPException(
                400,
                "Word table contains an English Translation column, but no "
                "non-empty English Translation segments were found.",
            )
        logger.info("Word: %d segments from table(s).", len(segments))
        return segments

    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if text:
            items.append(("", text))

    segments = _number_word_segments(items)
    if not segments:
        raise HTTPException(
            400,
            "Word document contains no usable interview text. Add interview "
            "paragraphs or a table with an 'English Translation' column.",
        )
    logger.info("Word: %d paragraph segments.", len(segments))
    return segments


# ---------------------------------------------------------------- dispatcher
def read_input_file(filename: str, raw: bytes) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return read_csv_file(raw)
    if name.endswith((".xlsx", ".xlsm")):
        return read_excel_file(raw)
    if name.endswith(".docx"):
        return read_word_file(raw)
    raise HTTPException(400, "Supported files: .csv, .xlsx, .xlsm, .docx")


# ---------------------------------------------------------------- output
HEADER_LABELS = {"MU_ID": "Segment_id"}  # internal column name -> header shown in Excel


def build_output_filename(uploaded_name: str, extension: str = "xlsx") -> str:
    """<uploaded file name>_PRISM_coded_<DD-MM-YYYY>_<hh-mm-ss>_<AM/PM>.<extension>"""
    stem = Path((uploaded_name or "").replace("\\", "/")).stem
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem).strip(" .")[:100] or "interview"
    now = datetime.now()
    hour12 = now.hour % 12 or 12
    am_pm = "AM" if now.hour < 12 else "PM"
    stamp = f"{now:%d-%m-%Y}_{hour12:02d}-{now:%M-%S}_{am_pm}"
    return f"{stem}_PRISM_coded_{stamp}.{extension}"

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
COLORED_COLUMNS = ("principal_code", "sub_code")


def _text_color_for(hex_color: str) -> str:
    """Black or white text, whichever is readable on this background."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return "000000" if (0.299 * r + 0.587 * g + 0.114 * b) > 160 else "FFFFFF"


def build_xlsx_bytes(output: list[dict[str, Any]]) -> bytes:
    """
    Excel output. A CSV cannot hold colors, so the result is an .xlsx where the
    principal_code and sub_code cells use the codebook color of their principal code.
    """
    if openpyxl is None:
        raise HTTPException(500, "openpyxl is not installed. Run: pip install openpyxl")
    if not output:
        raise HTTPException(400, "No output rows.")

    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    fieldnames: list[str] = []
    for row in output:
        for key in row:
            if key not in fieldnames and not key.startswith("_"):  # "_" keys are internal
                fieldnames.append(key)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "PRISM coding"

    # Show "Segment_id" instead of "MU_ID" (unless the input already has such a column).
    headers = [
        HEADER_LABELS[f] if f in HEADER_LABELS and HEADER_LABELS[f] not in fieldnames else f
        for f in fieldnames
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True)

    colored_idx = [i for i, f in enumerate(fieldnames) if f in COLORED_COLUMNS]
    speaker_field = find_field(fieldnames, SPEAKER_NAMES)
    speaker_idx = fieldnames.index(speaker_field) if speaker_field else None
    bold_font = Font(name="Calibri", size=11, bold=True)
    principal_idx = fieldnames.index("principal_code") if "principal_code" in fieldnames else None
    style_cache: dict[str, tuple[PatternFill, Font]] = {}
    widths = [len(h) for h in headers]

    for row in output:
        values = []
        for i, field in enumerate(fieldnames):
            value = row.get(field, "")
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            values.append(value)
            widths[i] = max(widths[i], min(len(str(value)), 60))
        sheet.append(values)

        excel_row = sheet[sheet.max_row]
        for cell in excel_row:
            # Interview text must never be interpreted as an Excel formula.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

        if speaker_idx is not None:
            excel_row[speaker_idx].font = bold_font

        principal = str(row.get("principal_code", "")).strip() if principal_idx is not None else ""
        color = PRINCIPAL_COLORS.get(principal)
        if color:
            if color not in style_cache:
                style_cache[color] = (
                    PatternFill(start_color=color, end_color=color, fill_type="solid"),
                    Font(name="Calibri", size=11, color=_text_color_for(color)),
                )
            fill, font = style_cache[color]
            for i in colored_idx:
                excel_row[i].fill = fill
                excel_row[i].font = font

    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width + 2
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- Word output
EVIDENCE_FILL = "D1E7DD"  # light green behind the evidence phrase
_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _shade(run, fill: str) -> None:
    """Background color behind a run of text."""
    run._r.get_or_add_rPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _evidence_spans(text: str, evidences: list[str]) -> list[tuple[int, int]]:
    """Where each evidence phrase sits inside the text (overlaps merged)."""
    spans = []
    for evidence in evidences:
        evidence = evidence.strip()
        match = re.search(re.escape(evidence), text, flags=re.IGNORECASE) if evidence else None
        if match:
            spans.append(match.span())

    merged: list[tuple[int, int]] = []
    for start, end in sorted(spans):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def _add_text_run(paragraph, text: str, highlight: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    if highlight:
        _shade(run, EVIDENCE_FILL)


def build_docx_bytes(output: list[dict[str, Any]]) -> bytes:
    """
    Word output: one block per segment - Segment ID + its codes on the left,
    interview text on the right with the evidence phrase(s) shaded light green.
    """
    if Document is None:
        raise HTTPException(500, "python-docx is not installed. Run: pip install python-docx")
    if not output:
        raise HTTPException(400, "No output rows.")

    # Group the one-row-per-code output back into one entry per segment.
    segments: dict[str, dict[str, Any]] = {}
    for row in output:
        segment_id = str(row.get("_segment_id", row.get("MU_ID", ""))).strip()
        entry = segments.setdefault(
            segment_id,
            {
                "text": _XML_ILLEGAL_RE.sub("", str(row.get("_text", ""))),
                "speaker": _XML_ILLEGAL_RE.sub("", str(row.get("_speaker", ""))).strip(),
                "codes": [],
            },
        )
        if str(row.get("principal_code", "")).strip():
            entry["codes"].append(row)

    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    title_p = doc.add_paragraph()
    run = title_p.add_run("PRISM AI Coded Interview Transcript")
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph()

    for segment_id, entry in segments.items():
        text = entry["text"]
        if not text:
            continue

        # 2-column borderless table: Segment ID + codes | interview text
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.2)
        table.columns[1].width = Inches(5.3)

        cell_meta, cell_text = table.rows[0].cells[0], table.rows[0].cells[1]
        cell_meta.width = Inches(1.2)
        cell_text.width = Inches(5.3)

        # Left cell: Segment ID, then each code (shaded with its codebook color)
        p_meta = cell_meta.paragraphs[0]
        r_id = p_meta.add_run(f"[{_XML_ILLEGAL_RE.sub('', segment_id)}]")
        r_id.font.name = "Calibri"
        r_id.font.size = Pt(9)
        r_id.font.bold = True
        r_id.font.color.rgb = RGBColor(100, 100, 100)

        for code in entry["codes"]:
            principal_code = str(code.get("principal_code", "")).strip()
            sub_code = str(code.get("sub_code", "")).strip()
            fill = PRINCIPAL_COLORS.get(principal_code)

            p_code = cell_meta.add_paragraph()
            p_code.paragraph_format.space_before = Pt(3)
            r_code = p_code.add_run(f"..{principal_code}")
            r_code.font.name = "Calibri"
            r_code.font.size = Pt(8.5)
            if fill:
                r_code.font.color.rgb = RGBColor.from_string(_text_color_for(fill))
                _shade(r_code, fill)
            else:
                r_code.font.color.rgb = RGBColor(0, 102, 204)

            if sub_code:
                r_sub = p_code.add_run(f"\n{sub_code}")
                r_sub.font.name = "Calibri"
                r_sub.font.size = Pt(8)
                r_sub.font.color.rgb = RGBColor(100, 100, 100)

        # Right cell: interview text, evidence phrases shaded light green
        p_txt = cell_text.paragraphs[0]
        p_txt.paragraph_format.line_spacing = 1.15

        if entry["speaker"]:
            r_speaker = p_txt.add_run(f"{entry['speaker']}: ")
            r_speaker.font.name = "Calibri"
            r_speaker.font.size = Pt(11)
            r_speaker.font.bold = True

        evidences = [_XML_ILLEGAL_RE.sub("", str(c.get("evidence", ""))) for c in entry["codes"]]
        position = 0
        for start, end in _evidence_spans(text, evidences):
            if start > position:
                _add_text_run(p_txt, text[position:start])
            _add_text_run(p_txt, text[start:end], highlight=True)
            position = end
        if position < len(text):
            _add_text_run(p_txt, text[position:])

        doc.add_paragraph()  # Spacer

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()